"""Small python-docx helpers for the KSC2026 experiment record (Korean fonts, tables, code blocks)."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import re as _re

_CTRL = _re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
KFONT = "맑은 고딕"
MONO = "Consolas"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
HEAD_FILL = "DCE8F7"
CODE_FILL = "F4F4F2"
NOTE_FILL = "FFF6E0"


def _set_fonts(rpr_parent, latin, east=KFONT):
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(k), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


class Report:
    def __init__(self):
        self.doc = Document()
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = d.styles["Normal"]
        st.font.name, st.font.size = KFONT, Pt(10)
        _set_fonts(st.element, KFONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size in ((1, 16), (2, 13), (3, 11)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = KFONT, Pt(size), True
            h.font.color.rgb = ACCENT if lvl < 3 else INK
            _set_fonts(h.element, KFONT)
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        t = d.styles["Title"]
        t.font.name, t.font.size = KFONT, Pt(24)
        _set_fonts(t.element, KFONT)
        code = d.styles.add_style("CodeBlock", 1)
        code.base_style = d.styles["Normal"]
        code.font.name, code.font.size = MONO, Pt(7.5)
        _set_fonts(code.element, MONO, KFONT)
        pf = code.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        cap = d.styles["Caption"]
        cap.font.name, cap.font.size, cap.font.italic = KFONT, Pt(9), False
        cap.font.color.rgb = INK2
        _set_fonts(cap.element, KFONT)
        self._footer_page_numbers()
        settings = d.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
        self.fig_no = 0
        self.tab_no = 0

    # ------------------------------------------------------------ text
    def title(self, text, sub=None):
        self.doc.add_paragraph(text, style="Title")
        if sub:
            p = self.doc.add_paragraph()
            r = p.add_run(sub)
            r.font.size, r.font.color.rgb = Pt(12), INK2

    def h(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def p(self, *parts, size=None, align=None, color=None):
        """parts: str or (str, 'b'|'i'|'c'|'bc') — b=bold, i=italic, c=inline code."""
        para = self.doc.add_paragraph()
        for part in parts:
            if isinstance(part, tuple):
                txt, fmt = part
            else:
                txt, fmt = part, ""
            r = para.add_run(txt)
            if "b" in fmt:
                r.bold = True
            if "i" in fmt:
                r.italic = True
            if "c" in fmt:
                r.font.name = MONO
                _set_fonts(r._element, MONO, KFONT)
                r.font.size = Pt(9)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return para

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.doc.add_paragraph(style=style)
            parts = it if isinstance(it, (list, tuple)) and not isinstance(it, str) else [it]
            for part in parts:
                if isinstance(part, tuple):
                    r = para.add_run(part[0])
                    if "b" in part[1]:
                        r.bold = True
                    if "c" in part[1]:
                        r.font.name = MONO
                        _set_fonts(r._element, MONO, KFONT)
                        r.font.size = Pt(9)
                else:
                    para.add_run(part)
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def note(self, text, fill=NOTE_FILL):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Cm(17)
        _shade(cell._tc.get_or_add_tcPr(), fill)
        lines = text if isinstance(text, list) else [text]
        cell.paragraphs[0].text = ""
        for i, line in enumerate(lines):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            r = para.add_run(line)
            r.font.size = Pt(9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def code(self, text, max_lines=None):
        lines = text.rstrip("\n").split("\n")
        if max_lines and len(lines) > max_lines:
            lines = lines[:max_lines] + [f"... ({len(lines) - max_lines} lines omitted)"]
        for line in lines:
            para = self.doc.add_paragraph(style="CodeBlock")
            _shade(para._p.get_or_add_pPr(), CODE_FILL)
            para.add_run(_CTRL.sub("", line.replace("\t", "    ")) if line else " ")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ table / figure
    def table(self, header, rows, widths=None, size=8.5, caption=None, align_right_from=None, bold_first_col=False):
        if caption:
            self.tab_no += 1
            cp = self.doc.add_paragraph(f"표 {self.tab_no}. {caption}", style="Caption")
            cp.paragraph_format.keep_with_next = True
        tbl = self.doc.add_table(rows=1, cols=len(header))
        tbl.style = self.doc.styles["Table Grid"]
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        hdr = tbl.rows[0]
        trpr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for i, txt in enumerate(header):
            c = hdr.cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(txt))
            r.bold, r.font.size = True, Pt(size)
            _shade(c._tc.get_or_add_tcPr(), HEAD_FILL)
        for row in rows:
            cells = tbl.add_row().cells
            for i, txt in enumerate(row):
                c = cells[i]
                c.text = ""
                lines = str(txt).split("\n")
                for k, line in enumerate(lines):
                    para = c.paragraphs[0] if k == 0 else c.add_paragraph()
                    para.paragraph_format.space_after = Pt(0)
                    r = para.add_run(line)
                    r.font.size = Pt(size)
                    if bold_first_col and i == 0:
                        r.bold = True
                    if align_right_from is not None and i >= align_right_from:
                        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths:
            for row in tbl.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return tbl

    def figure(self, path, caption, width_cm=16.5):
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        self.fig_no += 1
        self.doc.add_paragraph(f"그림 {self.fig_no}. {caption}", style="Caption").alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------------------------------------ fields
    def toc(self):
        para = self.doc.add_paragraph()
        r = para.add_run()
        for kind, text in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                           ("text", "목차: Word 에서 열 때 '필드 업데이트'를 누르면 채워진다."), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
                r._r.append(el)
            elif kind == "text":
                r2 = para.add_run(text)
                r2.font.color.rgb = INK2
                r = para.add_run()
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                r._r.append(el)

    def _footer_page_numbers(self):
        footer = self.doc.sections[0].footer
        para = footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run()
        for kind in ("begin", "instr", "end"):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = "PAGE"
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        r.font.size = Pt(9)

    def save(self, path):
        self.doc.save(str(path))
